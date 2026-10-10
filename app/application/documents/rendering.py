"""Render package artifact content as real, formatted DOCX and PDF documents."""
from __future__ import annotations

from io import BytesIO
from typing import Any
from xml.sax.saxutils import escape

from docx import Document as DocxDocument
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


MIME_TYPES = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
}
ARTIFACT_TITLES = {
    "tailored_cv": "Tailored CV",
    "cover_letter": "Cover Letter",
    "application_q_and_a": "Application Questions and Answers",
}


def _blocks(content: dict[str, Any], artifact_type: str) -> list[tuple[str, str]]:
    if not isinstance(content, dict):
        raise ValueError("document content must be an object")
    blocks: list[tuple[str, str]] = []
    title = content.get("title") or content.get("name") or ARTIFACT_TITLES.get(artifact_type, "Application Document")
    if isinstance(title, str) and title.strip():
        blocks.append(("title", title.strip()))

    for key in ("rawText", "content", "text", "body"):
        value = content.get(key)
        if isinstance(value, str) and value.strip():
            if not blocks:
                blocks.append(("title", ARTIFACT_TITLES.get(artifact_type, "Application Document")))
            for raw_line in value.splitlines():
                line = raw_line.strip()
                if not line:
                    continue
                if line.startswith("### "):
                    blocks.append(("heading2", line[4:].strip()))
                elif line.startswith("## "):
                    blocks.append(("heading", line[3:].strip()))
                elif line.startswith("# "):
                    blocks.append(("heading", line[2:].strip()))
                elif line.startswith(("- ", "* ")):
                    blocks.append(("bullet", line[2:].strip()))
                else:
                    blocks.append(("paragraph", line))
            return blocks

    answers = content.get("answers")
    if isinstance(answers, list):
        for index, answer in enumerate(answers, start=1):
            if not isinstance(answer, dict):
                continue
            question = answer.get("question")
            response = answer.get("answer")
            if isinstance(question, str) and question.strip():
                blocks.append(("heading", question.strip()))
            if isinstance(response, str) and response.strip():
                blocks.append(("paragraph", response.strip()))
        if len(blocks) > (1 if title else 0):
            return blocks

    sections = content.get("sections")
    if isinstance(sections, list):
        for section in sections:
            if not isinstance(section, dict):
                continue
            heading = section.get("heading") or section.get("title") or section.get("name")
            if isinstance(heading, str) and heading.strip():
                blocks.append(("heading", heading.strip()))
            body = section.get("body") or section.get("content") or section.get("text")
            if isinstance(body, str) and body.strip():
                blocks.extend(("paragraph", line.strip()) for line in body.splitlines() if line.strip())
            items = section.get("items") or section.get("bullets")
            if isinstance(items, list):
                blocks.extend(("bullet", item.strip()) for item in items if isinstance(item, str) and item.strip())
        if len(blocks) > (1 if title else 0):
            return blocks

    if len(blocks) <= 1:
        blocks = []
        for key, value in content.items():
            if key in {"title", "name", "rawText", "content", "text", "body", "metadata"}:
                continue
            if isinstance(value, str) and value.strip():
                blocks.append(("heading", key.replace("_", " ").title()))
                blocks.append(("paragraph", value.strip()))
            elif isinstance(value, list) and value:
                blocks.append(("heading", key.replace("_", " ").title()))
                for item in value:
                    if isinstance(item, str) and item.strip():
                        blocks.append(("bullet", item.strip()))
                    elif isinstance(item, dict):
                        for subkey, subvalue in item.items():
                            if isinstance(subvalue, str) and subvalue.strip():
                                blocks.append(("paragraph", f"{subkey.replace('_', ' ').title()}: {subvalue.strip()}"))
    if not blocks:
        raise ValueError("document content contains no renderable text")
    return blocks


def render_document(content: dict[str, Any], artifact_type: str, output_format: str) -> tuple[bytes, str, str]:
    """Return (bytes, MIME type, extension) without persisting the document."""
    fmt = output_format.lower() if isinstance(output_format, str) else ""
    if fmt not in MIME_TYPES:
        raise ValueError("output_format must be docx or pdf")
    blocks = _blocks(content, artifact_type)
    if fmt == "docx":
        output = BytesIO()
        document = DocxDocument()
        section = document.sections[0]
        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        normal = document.styles["Normal"]
        normal.font.name = "Arial"
        normal.font.size = Pt(10)
        for kind, text in blocks:
            if kind == "title":
                document.add_heading(text, level=0)
            elif kind == "heading":
                document.add_heading(text, level=1)
            elif kind == "heading2":
                document.add_heading(text, level=2)
            elif kind == "bullet":
                document.add_paragraph(text, style="List Bullet")
            else:
                document.add_paragraph(text)
        document.save(output)
        return output.getvalue(), MIME_TYPES[fmt], fmt

    output = BytesIO()
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name="DocumentTitle", parent=styles["Title"], fontName="Helvetica-Bold",
        fontSize=19, leading=23, textColor=colors.HexColor("#18243A"), spaceAfter=18,
        alignment=TA_LEFT,
    ))
    styles.add(ParagraphStyle(
        name="DocumentHeading", parent=styles["Heading2"], fontName="Helvetica-Bold",
        fontSize=12, leading=15, textColor=colors.HexColor("#243B64"), spaceBefore=10,
        spaceAfter=5, keepWithNext=True,
    ))
    styles.add(ParagraphStyle(
        name="DocumentBody", parent=styles["BodyText"], fontName="Helvetica",
        fontSize=9.5, leading=13, spaceAfter=7,
    ))
    story = []
    for kind, text in blocks:
        safe_text = escape(text).replace("\n", "<br/>")
        if kind == "title":
            story.append(Paragraph(safe_text, styles["DocumentTitle"]))
        elif kind in {"heading", "heading2"}:
            story.append(Paragraph(safe_text, styles["DocumentHeading"]))
        elif kind == "bullet":
            story.append(Paragraph(f"- {safe_text}", styles["DocumentBody"]))
        else:
            story.append(Paragraph(safe_text, styles["DocumentBody"]))
    SimpleDocTemplate(
        output, pagesize=LETTER, rightMargin=0.7 * inch, leftMargin=0.7 * inch,
        topMargin=0.65 * inch, bottomMargin=0.65 * inch,
        title=ARTIFACT_TITLES.get(artifact_type, "Application Document"),
        author="JobPunter",
    ).build(story)
    return output.getvalue(), MIME_TYPES[fmt], fmt
