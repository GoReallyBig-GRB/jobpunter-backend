"""Validated candidate-private document import, extraction, and preview."""
from __future__ import annotations

import json
import logging
import re
from datetime import datetime, timezone
import zipfile
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from uuid import UUID, uuid4

from docx import Document as DocxDocument
from pypdf import PdfReader
from fastapi import HTTPException
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.api.errors import DomainError
from app.domain.candidate.provenance import source_excerpt_is_verifiable
from app.application.ai.schemas import GenerateOptions, GenerateRequest
from app.application.ai.service import execute_ai_task
from app.application.documents.structured_conflicts import conflict_resolution_key, find_structured_conflicts
from app.application.documents.backup import backup_to_personal_drive
from app.config.settings import get_settings
from app.infrastructure.db.models import CandidateProfile, CandidateProfileFactProvenance, Document, DocumentImport, Evidence, EvidenceSource
from app.infrastructure.storage.factory import build_object_storage
from app.infrastructure.ocr.google_vision import OCRProviderError, OCRUnavailableError, ocr_pdf_pages
from app.ports.object_storage import ObjectStorage, StoredObject

logger = logging.getLogger("jobpunter.document_imports")
MAX_DOCUMENT_BYTES = 4 * 1024 * 1024
MAX_DOCX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_PDF_PAGES = 200
MIME_BY_EXTENSION = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".txt": "text/plain",
}
GENERIC_MIME_TYPES = {"application/octet-stream", "binary/octet-stream"}


def validate_document_upload(filename: str, content: bytes, content_type: str) -> tuple[str, str]:
    if not isinstance(content, bytes) or not content:
        raise DomainError("DOCUMENT_EMPTY", "Uploaded document is empty.", 422)
    if len(content) > MAX_DOCUMENT_BYTES:
        raise DomainError("DOCUMENT_TOO_LARGE", "Uploaded document exceeds the 4 MB limit.", 413)
    extension = Path(filename or "").suffix.lower()
    mime_type = MIME_BY_EXTENSION.get(extension)
    if mime_type is None:
        raise DomainError(
            "DOCUMENT_FORMAT_UNSUPPORTED",
            "Only PDF, DOCX, and TXT files are supported.",
            415,
        )
    supplied_type = (content_type or "").split(";", 1)[0].strip().lower()
    if supplied_type and supplied_type not in GENERIC_MIME_TYPES | {mime_type}:
        raise DomainError(
            "DOCUMENT_CONTENT_TYPE_MISMATCH",
            "Uploaded content type does not match the file extension.",
            415,
        )
    if extension == ".pdf":
        if not content.startswith(b"%PDF-"):
            raise DomainError("DOCUMENT_SIGNATURE_INVALID", "File is not a valid PDF upload.", 422)
    elif extension == ".docx":
        if not zipfile.is_zipfile(BytesIO(content)):
            raise DomainError("DOCUMENT_SIGNATURE_INVALID", "File is not a valid DOCX upload.", 422)
        try:
            with zipfile.ZipFile(BytesIO(content)) as archive:
                names = set(archive.namelist())
                if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                    raise DomainError("DOCUMENT_SIGNATURE_INVALID", "File is not a valid DOCX document.", 422)
                if sum(info.file_size for info in archive.infolist()) > MAX_DOCX_UNCOMPRESSED_BYTES:
                    raise DomainError("DOCUMENT_TOO_LARGE", "DOCX expands beyond the safe processing limit.", 413)
        except zipfile.BadZipFile as exc:
            raise DomainError("DOCUMENT_SIGNATURE_INVALID", "File is not a valid DOCX upload.", 422) from exc
    else:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise DomainError("DOCUMENT_ENCODING_UNSUPPORTED", "TXT uploads must use UTF-8 encoding.", 422) from exc
        if "\x00" in text:
            raise DomainError("DOCUMENT_ENCODING_UNSUPPORTED", "TXT uploads contain unsupported binary data.", 422)
    return extension, mime_type


def extract_document_text(extension: str, content: bytes) -> str:
    try:
        if extension == ".txt":
            text = content.decode("utf-8-sig")
        elif extension == ".docx":
            document = DocxDocument(BytesIO(content))
            parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
            text = "\n".join(parts)
        elif extension == ".pdf":
            reader = PdfReader(BytesIO(content), strict=True)
            if reader.is_encrypted:
                raise DomainError("DOCUMENT_ENCRYPTED", "Password-protected PDFs are not supported.", 422)
            if len(reader.pages) > MAX_PDF_PAGES:
                raise DomainError("DOCUMENT_PAGE_LIMIT_EXCEEDED", "PDF exceeds the 200-page processing limit.", 413)
            page_texts = []
            for page in reader.pages:
                try:
                    page_texts.append(page.extract_text() or "")
                except Exception:
                    # A page without selectable text may still be readable by OCR.
                    page_texts.append("")
            scanned_pages = [index for index, page_text in enumerate(page_texts) if not page_text.strip()]
            if scanned_pages:
                settings = get_settings()
                if not settings.document_ocr_enabled:
                    raise DomainError(
                        "DOCUMENT_OCR_DISABLED",
                        "This PDF contains scanned pages. Enable DOCUMENT_OCR_ENABLED after authorizing the OCR provider.",
                        503,
                    )
                try:
                    ocr_text_by_page = ocr_pdf_pages(
                        content,
                        scanned_pages,
                        timeout_seconds=settings.document_ocr_timeout_seconds,
                        batch_size=settings.document_ocr_batch_size,
                    )
                except OCRUnavailableError as exc:
                    raise DomainError(
                        "DOCUMENT_OCR_UNAVAILABLE",
                        "Scanned-PDF OCR is enabled but the OCR provider is not available.",
                        503,
                    ) from exc
                except OCRProviderError as exc:
                    raise DomainError(
                        "DOCUMENT_OCR_FAILED",
                        "Scanned-PDF OCR failed. Retry the import after checking the OCR provider.",
                        502,
                    ) from exc
                for page_index in scanned_pages:
                    ocr_text = ocr_text_by_page.get(page_index, "").strip()
                    if not ocr_text:
                        raise DomainError(
                            "DOCUMENT_OCR_TEXT_NOT_FOUND",
                            f"OCR did not find readable text on PDF page {page_index + 1}.",
                            422,
                        )
                    page_texts[page_index] = ocr_text
            text = "\n".join(page_texts)
        else:
            raise DomainError("DOCUMENT_FORMAT_UNSUPPORTED", "Only PDF, DOCX, and TXT files are supported.", 415)
    except DomainError:
        raise
    except Exception as exc:
        raise DomainError("DOCUMENT_PARSE_FAILED", "The uploaded document could not be parsed.", 422) from exc
    text = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    if not text:
        raise DomainError(
            "DOCUMENT_TEXT_NOT_EXTRACTABLE",
            "No readable text was found in the uploaded document.",
            422,
        )
    return text


def _safe_filename(filename: str, extension: str) -> str:
    name = Path(filename or "").name.replace("\x00", "").strip()
    name = re.sub(r"[^A-Za-z0-9._ -]", "_", name)
    if not name or name in {".", ".."}:
        name = f"uploaded-document{extension}"
    return name[:255]


class DocumentImportService:
    def __init__(self, db: Session, storage: ObjectStorage | None = None):
        self.db = db
        self._storage = storage

    def _profile(self, user_id: UUID) -> CandidateProfile:
        profile = self.db.scalar(select(CandidateProfile).where(CandidateProfile.user_id == user_id))
        if profile is None:
            raise DomainError("CANDIDATE_PROFILE_NOT_FOUND", "Candidate profile not found", 404)
        return profile

    def _storage_adapter(self) -> ObjectStorage:
        if self._storage is not None:
            return self._storage
        try:
            self._storage = build_object_storage(get_settings())
        except DomainError:
            raise
        except Exception as exc:
            logger.warning("document import storage initialization failed: %s", type(exc).__name__)
            raise DomainError("OBJECT_STORAGE_UNAVAILABLE", "Private document storage is unavailable.", 503) from exc
        return self._storage

    def upload(
        self, user_id: UUID, filename: str, content_type: str, content: bytes
    ) -> DocumentImport:
        profile = self._profile(user_id)
        extension, mime_type = validate_document_upload(filename, content, content_type)
        safe_name = _safe_filename(filename, extension)
        document_id = uuid4()
        storage_key = f"candidates/{profile.id}/imports/{document_id}/{safe_name}"
        storage = self._storage_adapter()
        checksum = sha256(content).hexdigest()
        try:
            stored: StoredObject = storage.put(storage_key, content, mime_type)
        except Exception as exc:
            logger.warning("document import upload failed: %s", type(exc).__name__)
            raise DomainError("DOCUMENT_STORAGE_FAILED", "Uploaded document could not be stored securely.", 502) from exc
        if (
            stored.storage_key != storage_key
            or stored.checksum_sha256 != checksum
            or stored.byte_size != len(content)
        ):
            try:
                storage.delete(storage_key)
            except Exception as exc:
                logger.error("failed to clean up invalid uploaded document error_type=%s", type(exc).__name__)
            raise DomainError("DOCUMENT_STORAGE_INTEGRITY_FAILED", "Stored document integrity could not be verified.", 502)

        drive_backup = backup_to_personal_drive(
            get_settings(),
            document_id=str(document_id),
            filename=safe_name,
            content=content,
            content_type=mime_type,
            checksum_sha256=checksum,
        )
        document = Document(
            id=document_id,
            candidate_profile_id=profile.id,
            document_type="import_source",
            filename=safe_name,
            mime_type=mime_type,
            storage_key=storage_key,
            checksum_sha256=checksum,
            byte_size=len(content),
            metadata_json={
                "original_content_type": content_type,
                "extension": extension,
                "google_drive_backup": drive_backup,
            },
        )
        record = DocumentImport(
            candidate_profile_id=profile.id,
            document_id=document_id,
            status="UPLOADED",
            extracted_text=None,
            extracted_fields={},
            conflicts=[],
            revision=1,
            error_code=None,
        )
        try:
            self.db.add(document)
            self.db.flush()
            self.db.add(record)
            self.db.commit()
            self.db.refresh(record)
        except Exception:
            self.db.rollback()
            try:
                storage.delete(storage_key)
            except Exception as exc:
                logger.error("failed to clean up uploaded object after metadata failure error_type=%s", type(exc).__name__)
            raise
        return record

    def _owned_import(self, user_id: UUID, import_id: UUID) -> tuple[CandidateProfile, DocumentImport, Document]:
        profile = self._profile(user_id)
        record = self.db.scalar(
            select(DocumentImport).where(
                DocumentImport.id == import_id,
                DocumentImport.candidate_profile_id == profile.id,
            )
            .with_for_update()
            .execution_options(populate_existing=True)
        )
        if record is None:
            raise DomainError("DOCUMENT_IMPORT_NOT_FOUND", "Document import not found", 404)
        document = self.db.scalar(select(Document).where(
            Document.id == record.document_id,
            Document.candidate_profile_id == profile.id,
        ))
        if document is None:
            raise DomainError("DOCUMENT_NOT_FOUND", "Imported document not found", 404)
        return profile, record, document

    def parse(self, user_id: UUID, import_id: UUID) -> DocumentImport:
        _profile, record, document = self._owned_import(user_id, import_id)
        if record.status in {"PARSED", "PREVIEWED", "CONFLICTS_IDENTIFIED", "CONFIRMED", "SAVED"}:
            return record
        if record.status == "FAILED":
            # A parse failure is recoverable while the original private object remains
            # available and intact. Reset only derived parse state; never create evidence.
            record.status = "UPLOADED"
            record.extracted_text = None
            record.extracted_fields = {}
            record.conflicts = []
            record.error_code = None
            record.revision += 1
            record.updated_at = datetime.now(timezone.utc)
            # _owned_import already holds this row's FOR UPDATE lock. Do not
            # commit the intermediate UPLOADED state: another retry could acquire
            # the row and race this parse before the lock is reacquired.
        if record.status != "UPLOADED":
            raise DomainError("DOCUMENT_IMPORT_STATE_CONFLICT", "This import cannot be parsed in its current state.", 409)
        try:
            content = self._storage_adapter().get(document.storage_key)
        except (FileNotFoundError, KeyError) as exc:
            record.status = "FAILED"
            record.error_code = "DOCUMENT_OBJECT_MISSING"
            record.revision += 1
            record.updated_at = datetime.now(timezone.utc)
            self.db.commit()
            raise DomainError(
                "DOCUMENT_OBJECT_MISSING",
                "The private document object is missing. Restore it or upload the document again.",
                409,
            ) from exc
        except DomainError as exc:
            record.status = "FAILED"
            record.error_code = exc.code
            record.revision += 1
            record.updated_at = datetime.now(timezone.utc)
            self.db.commit()
            raise
        except Exception as exc:
            record.status = "FAILED"
            record.error_code = "DOCUMENT_STORAGE_READ_FAILED"
            record.revision += 1
            record.updated_at = datetime.now(timezone.utc)
            self.db.commit()
            logger.warning("document import storage read failed: %s", type(exc).__name__)
            raise DomainError(
                "DOCUMENT_STORAGE_READ_FAILED",
                "Private document storage could not be read. Retry parsing after storage is available.",
                503,
            ) from exc

        try:
            if len(content) != document.byte_size or sha256(content).hexdigest() != document.checksum_sha256:
                raise DomainError("DOCUMENT_INTEGRITY_FAILED", "Uploaded document failed integrity verification.", 409)
            extension = document.metadata_json.get("extension")
            record.extracted_text = extract_document_text(extension, content)
            record.status = "PARSED"
            record.error_code = None
        except DomainError as exc:
            record.status = "FAILED"
            record.error_code = exc.code
            record.revision += 1
            record.updated_at = datetime.now(timezone.utc)
            self.db.commit()
            raise
        except Exception as exc:
            record.status = "FAILED"
            record.error_code = "DOCUMENT_PARSE_FAILED"
            record.revision += 1
            record.updated_at = datetime.now(timezone.utc)
            self.db.commit()
            logger.warning("document import parse failed: %s", type(exc).__name__)
            raise DomainError("DOCUMENT_PARSE_FAILED", "The uploaded document could not be parsed.", 422) from exc
        record.revision += 1
        self.db.commit()
        self.db.refresh(record)
        return record


    def _record_extraction_failure(
        self,
        import_id: UUID,
        *,
        expected_revision: int,
        expected_status: str,
        error_code: str,
    ) -> None:
        current = self.db.scalar(
            select(DocumentImport)
            .where(DocumentImport.id == import_id)
            .with_for_update()
            .execution_options(populate_existing=True)
        )
        if current is None:
            raise DomainError("DOCUMENT_IMPORT_NOT_FOUND", "Document import no longer exists.", 404)
        if current.revision != expected_revision or current.status != expected_status:
            raise DomainError(
                "REVISION_CONFLICT",
                "Document import changed while extraction was running. Refresh the import and retry.",
                409,
                {"current_revision": current.revision},
            )
        current.error_code = error_code
        current.revision += 1
        current.updated_at = datetime.now(timezone.utc)
        self.db.commit()

    async def extract_fields(self, user_id: UUID, import_id: UUID) -> DocumentImport:
        profile, record, _document = self._owned_import(user_id, import_id)
        if record.status not in {"PARSED", "PREVIEWED"} or not record.extracted_text:
            raise DomainError(
                "DOCUMENT_IMPORT_NOT_EXTRACTABLE",
                "Parse the document before requesting structured extraction.",
                409,
            )
        # The server-owned document_parsing contract and provider input limits make
        # truncation unsafe. Ask the candidate to split unusually long documents.
        if len(record.extracted_text) > 70_000:
            raise DomainError(
                "DOCUMENT_TEXT_TOO_LONG_FOR_EXTRACTION",
                "This document contains too much text for one safe extraction request. Split it into smaller documents.",
                413,
            )
        expected_revision = record.revision
        expected_status = record.status
        request = GenerateRequest(
            task="document_parsing",
            prompt=json.dumps(
                {"document_text": record.extracted_text},
                ensure_ascii=False,
                separators=(",", ":"),
            ),
            options=GenerateOptions(jsonMode=True, maxTokens=4096),
        )
        try:
            response = await execute_ai_task(
                request, db=self.db, user_id=profile.user_id
            )
        except HTTPException as exc:
            self.db.rollback()
            detail = exc.detail
            reported_code = detail.get("code") if isinstance(detail, dict) else None
            error_code = (
                reported_code
                if isinstance(reported_code, str) and re.fullmatch(r"[A-Z0-9_]{1,64}", reported_code)
                else "DOCUMENT_EXTRACTION_FAILED"
            )
            self._record_extraction_failure(
                record.id,
                expected_revision=expected_revision,
                expected_status=expected_status,
                error_code=error_code,
            )
            raise
        except Exception as exc:
            self.db.rollback()
            logger.warning("document import extraction failed: %s", type(exc).__name__)
            self._record_extraction_failure(
                record.id,
                expected_revision=expected_revision,
                expected_status=expected_status,
                error_code="DOCUMENT_EXTRACTION_FAILED",
            )
            raise DomainError(
                "DOCUMENT_EXTRACTION_FAILED",
                "Structured extraction failed. The parsed document remains available for retry.",
                502,
            ) from exc
        if response.execution_id is None:
            self._record_extraction_failure(
                record.id,
                expected_revision=expected_revision,
                expected_status=expected_status,
                error_code="AI_EXECUTION_TRACE_MISSING",
            )
            raise DomainError(
                "AI_EXECUTION_TRACE_MISSING",
                "Structured extraction completed without a durable AI execution reference.",
                502,
            )
        # AI execution commits its telemetry and releases the import-row lock.
        # Re-lock and verify the revision before applying the slower extraction result,
        # otherwise a concurrent candidate confirmation could be overwritten.
        record = self.db.scalar(
            select(DocumentImport)
            .where(DocumentImport.id == record.id)
            .with_for_update()
            .execution_options(populate_existing=True)
        )
        if record is None:
            raise DomainError(
                "DOCUMENT_IMPORT_NOT_FOUND",
                "Document import no longer exists. Refresh and retry.",
                404,
            )
        if record.revision != expected_revision or record.status != expected_status:
            raise DomainError(
                "REVISION_CONFLICT",
                "Document import changed while extraction was running. Refresh the import and retry.",
                409,
                {"current_revision": record.revision},
            )
        parsed = response.parsed_data
        if not isinstance(parsed, dict) or not isinstance(parsed.get("fields"), list):
            record.error_code = "DOCUMENT_EXTRACTION_OUTPUT_INVALID"
            record.revision += 1
            record.updated_at = datetime.now(timezone.utc)
            self.db.commit()
            raise DomainError(
                "DOCUMENT_EXTRACTION_OUTPUT_INVALID",
                "Structured extraction did not return a valid document result. Please retry.",
                502,
            )
        record.ai_execution_id = response.execution_id
        record.error_code = None
        record.extracted_fields = parsed
        # Model-reported conflicts are preserved as review signals. They never
        # reject evidence or bypass candidate confirmation.
        record.conflicts = [
            {"code": "EXTRACTION_REPORTED_CONFLICT", "message": item}
            for item in parsed.get("conflicts", [])
            if isinstance(item, str) and item.strip()
        ]
        record.status = "PREVIEWED"
        record.revision += 1
        record.updated_at = datetime.now(timezone.utc)
        self.db.commit()
        self.db.refresh(record)
        return record

    def identify_conflicts(self, user_id: UUID, import_id: UUID) -> DocumentImport:
        profile, record, _document = self._owned_import(user_id, import_id)
        if record.status == "CONFLICTS_IDENTIFIED":
            return record
        if record.status not in {"PARSED", "PREVIEWED"} or not record.extracted_text:
            raise DomainError(
                "DOCUMENT_IMPORT_NOT_REVIEWABLE",
                "Parse and preview the document before identifying conflicts.",
                409,
            )
        existing = self.db.scalars(select(Evidence).where(
            Evidence.candidate_profile_id == profile.id,
            Evidence.status == "APPROVED",
        ).order_by(Evidence.id)).all()
        conflicts = list(record.conflicts or [])
        if isinstance(record.extracted_fields, dict) and record.extracted_fields.get("fields"):
            conflicts.extend(find_structured_conflicts(record.extracted_fields, existing, profile.profile_data))
        else:
            # Compatibility fallback for imports not yet processed by structured
            # extraction. Overlap is only a review signal, not a semantic verdict.
            normalized_text = " ".join(record.extracted_text.casefold().split())
            for evidence in existing:
                for value in _flatten_strings(evidence.content):
                    normalized_value = " ".join(value.casefold().split())
                    if len(normalized_value) >= 12 and normalized_value in normalized_text:
                        conflicts.append({
                            "code": "POSSIBLE_EVIDENCE_OVERLAP",
                            "evidence_id": str(evidence.id),
                            "evidence_type": evidence.evidence_type,
                            "message": "Imported text overlaps existing evidence; review both records before saving.",
                        })
                        break
        # Stable de-duplication avoids repeated model or field signals on retries.
        deduplicated = []
        seen_conflicts = set()
        for conflict in conflicts:
            key = (
                conflict.get("code"),
                conflict.get("field"),
                conflict.get("evidence_id"),
                str(conflict.get("imported_value") or conflict.get("imported_values") or ""),
                str(conflict.get("existing_value") or ""),
                conflict.get("message"),
            )
            if key not in seen_conflicts:
                seen_conflicts.add(key)
                deduplicated.append(conflict)
        record.conflicts = deduplicated
        record.conflict_resolutions = []
        record.status = "CONFLICTS_IDENTIFIED"
        record.revision += 1
        record.updated_at = datetime.now(timezone.utc)
        self.db.commit()
        self.db.refresh(record)
        return record

    def confirm(
        self,
        user_id: UUID,
        import_id: UUID,
        *,
        expected_revision: int,
        conflicts_reviewed: bool = False,
        conflict_resolutions: list[dict] | None = None,
    ) -> DocumentImport:
        profile, record, _document = self._owned_import(user_id, import_id)
        if record.revision != expected_revision:
            raise DomainError(
                "REVISION_CONFLICT",
                "Document import revision is stale.",
                409,
                {"current_revision": record.revision},
            )
        if record.status != "CONFLICTS_IDENTIFIED":
            raise DomainError(
                "DOCUMENT_IMPORT_STATE_CONFLICT",
                "Conflict identification must finish before candidate confirmation.",
                409,
            )
        conflicts = record.conflicts if isinstance(record.conflicts, list) else []
        resolutions = conflict_resolutions or []
        if conflicts and not conflicts_reviewed:
            raise DomainError(
                "DOCUMENT_CONFLICT_REVIEW_REQUIRED",
                "Review the possible overlaps before confirming this import.",
                409,
                {"conflicts": conflicts},
            )
        if conflicts:
            expected_keys = [conflict_resolution_key(conflict) for conflict in conflicts if isinstance(conflict, dict)]
            if len(expected_keys) != len(conflicts):
                raise DomainError(
                    "DOCUMENT_CONFLICT_DATA_INVALID",
                    "Stored conflict data cannot be resolved safely.",
                    409,
                )
            decision_by_key = {}
            allowed_decisions = {"KEEP_IMPORTED", "KEEP_EXISTING", "KEEP_BOTH"}
            for resolution in resolutions:
                if not isinstance(resolution, dict):
                    raise DomainError(
                        "DOCUMENT_CONFLICT_RESOLUTIONS_INVALID",
                        "Each conflict requires a valid candidate decision.",
                        422,
                    )
                key = resolution.get("conflict_key")
                decision = resolution.get("decision")
                if (
                    not isinstance(key, str)
                    or not isinstance(decision, str)
                    or key in decision_by_key
                    or key not in expected_keys
                    or decision not in allowed_decisions
                ):
                    raise DomainError(
                        "DOCUMENT_CONFLICT_RESOLUTIONS_INVALID",
                        "Conflict decisions must be unique, current, and use a supported decision.",
                        422,
                        {"expected_conflict_keys": expected_keys},
                    )
                decision_by_key[key] = decision
            if set(decision_by_key) != set(expected_keys) or len(decision_by_key) != len(expected_keys):
                raise DomainError(
                    "DOCUMENT_CONFLICT_RESOLUTIONS_REQUIRED",
                    "Resolve every displayed conflict before confirming this import.",
                    409,
                    {"expected_conflict_keys": expected_keys},
                )
            record.conflict_resolutions = [
                {"conflict_key": key, "decision": decision_by_key[key]}
                for key in expected_keys
            ]
        else:
            if resolutions:
                raise DomainError(
                    "DOCUMENT_CONFLICT_RESOLUTIONS_INVALID",
                    "This import has no conflicts to resolve.",
                    422,
                )
            record.conflict_resolutions = []
        record.status = "CONFIRMED"
        record.confirmed_at = datetime.now(timezone.utc)
        record.confirmed_by_user_id = profile.user_id
        record.revision += 1
        record.updated_at = datetime.now(timezone.utc)
        self.db.commit()
        self.db.refresh(record)
        return record

    def save(self, user_id: UUID, import_id: UUID) -> tuple[DocumentImport, Evidence]:
        profile, record, document = self._owned_import(user_id, import_id)
        if record.status == "SAVED" and record.saved_evidence_id:
            evidence = self.db.get(Evidence, record.saved_evidence_id)
            if evidence is not None and evidence.candidate_profile_id == profile.id:
                return record, evidence
            raise DomainError(
                "DOCUMENT_IMPORT_SAVE_INTEGRITY_FAILED",
                "Saved import evidence could not be verified.",
                409,
            )
        if record.conflicts:
            expected_keys = {
                conflict_resolution_key(conflict)
                for conflict in record.conflicts if isinstance(conflict, dict)
            }
            saved_keys = {
                item.get("conflict_key") for item in (record.conflict_resolutions or [])
                if isinstance(item, dict)
            }
            if expected_keys != saved_keys:
                raise DomainError(
                    "DOCUMENT_CONFLICT_RESOLUTIONS_REQUIRED",
                    "Resolve every displayed conflict before saving this import.",
                    409,
                )
        if record.status != "CONFIRMED" or not record.confirmed_at or record.confirmed_by_user_id != profile.user_id:
            raise DomainError(
                "DOCUMENT_IMPORT_CONFIRMATION_REQUIRED",
                "Candidate confirmation is required before saving extracted information.",
                409,
            )
        source = EvidenceSource(
            source_type="candidate_upload",
            source_reference=f"document-import:{record.id}",
            metadata_json={
                "document_id": str(document.id),
                "document_import_id": str(record.id),
                "filename": document.filename,
                "checksum_sha256": document.checksum_sha256,
            },
        )
        try:
            self.db.add(source)
            self.db.flush()
            evidence = Evidence(
                candidate_profile_id=profile.id,
                evidence_source_id=source.id,
                status="DRAFT",
                evidence_type="imported_document",
                content={
                    "extracted_text": record.extracted_text or "",
                    "extracted_fields": record.extracted_fields or {},
                    "document_id": str(document.id),
                    "document_import_id": str(record.id),
                    "filename": document.filename,
                },
            )
            self.db.add(evidence)
            self.db.flush()
            extracted_fields = record.extracted_fields if isinstance(record.extracted_fields, dict) else {}
            raw_facts = extracted_fields.get("fields", [])
            if isinstance(raw_facts, list):
                for index, field in enumerate(raw_facts):
                    if not isinstance(field, dict) or "value" not in field or field["value"] is None:
                        continue
                    field_name = str(field.get("name") or "unknown").strip()[:255] or "unknown"
                    candidate_excerpt = field.get("source_text")
                    source_excerpt = (
                        candidate_excerpt
                        if source_excerpt_is_verifiable(record.extracted_text, candidate_excerpt)
                        else None
                    )
                    self.db.add(CandidateProfileFactProvenance(
                        candidate_profile_id=profile.id,
                        section="imported_document",
                        fact_path=f"fields[{index}].{field_name}",
                        fact_value=field["value"],
                        profile_revision=profile.revision,
                        source_type="DOCUMENT_IMPORT",
                        source_reference=f"document-import:{record.id}",
                        evidence_id=evidence.id,
                        document_import_id=record.id,
                        source_excerpt=source_excerpt if isinstance(source_excerpt, str) else None,
                        confidence=field.get("confidence"),
                        confirmed_at=record.confirmed_at,
                        confirmed_by_user_id=record.confirmed_by_user_id,
                        created_at=datetime.now(timezone.utc),
                    ))
            record.saved_evidence_id = evidence.id
            record.status = "SAVED"
            record.revision += 1
            record.updated_at = datetime.now(timezone.utc)
            self.db.commit()
            self.db.refresh(record)
            self.db.refresh(evidence)
            return record, evidence
        except Exception:
            self.db.rollback()
            raise

    def preview(self, user_id: UUID, import_id: UUID) -> DocumentImport:
        _profile, record, _document = self._owned_import(user_id, import_id)
        if record.status not in {"PARSED", "PREVIEWED", "CONFLICTS_IDENTIFIED"}:
            raise DomainError("DOCUMENT_IMPORT_NOT_PREVIEWABLE", "Parse the document before requesting its preview.", 409)
        if record.status == "PARSED":
            record.status = "PREVIEWED"
            record.revision += 1
            self.db.commit()
            self.db.refresh(record)
        return record


    def list_history(
        self,
        user_id: UUID,
        *,
        limit: int = 50,
        offset: int = 0,
        status: str | None = None,
    ) -> tuple[list[tuple[DocumentImport, Document]], int]:
        """List candidate-owned import metadata without returning extracted document contents."""
        profile = self._profile(user_id)
        filters = [DocumentImport.candidate_profile_id == profile.id]
        if status is not None:
            filters.append(DocumentImport.status == status)
        total = self.db.scalar(
            select(func.count()).select_from(DocumentImport).where(*filters)
        ) or 0
        records = self.db.execute(
            select(DocumentImport, Document)
            .join(
                Document,
                (Document.id == DocumentImport.document_id)
                & (Document.candidate_profile_id == profile.id),
            )
            .where(*filters)
            .order_by(DocumentImport.created_at.desc(), DocumentImport.id.desc())
            .offset(offset)
            .limit(limit)
        ).all()
        return [(record, document) for record, document in records], total

    def get_history_item(
        self, user_id: UUID, import_id: UUID
    ) -> tuple[DocumentImport, Document]:
        """Return one candidate-owned import and its metadata, never the stored file bytes."""
        _profile, record, document = self._owned_import(user_id, import_id)
        return record, document

def _flatten_strings(value) -> list[str]:
    if isinstance(value, str):
        return [value]
    if isinstance(value, dict):
        return [text for child in value.values() for text in _flatten_strings(child)]
    if isinstance(value, list):
        return [text for child in value for text in _flatten_strings(child)]
    return []

