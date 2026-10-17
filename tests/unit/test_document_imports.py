from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4

import pytest
from docx import Document as DocxDocument
from reportlab.pdfgen import canvas

from app.api.errors import DomainError
from app.application.documents.imports import extract_document_text, validate_document_upload


def _docx_bytes():
    output = BytesIO()
    document = DocxDocument()
    document.add_heading("Career History", level=1)
    document.add_paragraph("Product Manager at Example Ltd.")
    document.save(output)
    return output.getvalue()


def _pdf_bytes():
    output = BytesIO()
    pdf = canvas.Canvas(output)
    pdf.drawString(72, 720, "Product Manager at Example Ltd.")
    pdf.save()
    return output.getvalue()


@pytest.mark.parametrize(
    ("filename", "content_type", "content", "extension", "mime_type"),
    [
        ("career.txt", "text/plain", b"Product Manager", ".txt", "text/plain"),
        ("career.docx", "application/octet-stream", _docx_bytes(), ".docx",
         "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        ("career.pdf", "application/pdf", _pdf_bytes(), ".pdf", "application/pdf"),
    ],
    ids=["txt", "docx", "pdf"],
)
def test_document_upload_validation_accepts_supported_formats(
    filename, content_type, content, extension, mime_type
):
    assert validate_document_upload(filename, content, content_type) == (extension, mime_type)


@pytest.mark.parametrize(
    ("filename", "content_type", "content", "code"),
    [
        ("career.exe", "application/octet-stream", b"not a document", "DOCUMENT_FORMAT_UNSUPPORTED"),
        ("career.pdf", "application/pdf", b"not a pdf", "DOCUMENT_SIGNATURE_INVALID"),
        ("career.txt", "text/plain", b"\xff\xfe", "DOCUMENT_ENCODING_UNSUPPORTED"),
        ("career.txt", "application/pdf", b"text", "DOCUMENT_CONTENT_TYPE_MISMATCH"),
        ("career.txt", "text/plain", b"", "DOCUMENT_EMPTY"),
    ],
)
def test_document_upload_validation_rejects_invalid_files(filename, content_type, content, code):
    with pytest.raises(DomainError) as error:
        validate_document_upload(filename, content, content_type)
    assert error.value.code == code


def test_document_text_extraction_supports_txt_docx_and_pdf():
    assert extract_document_text(".txt", b"Plain text\r\nSecond line") == "Plain text\nSecond line"
    assert "Product Manager" in extract_document_text(".docx", _docx_bytes())
    assert "Example Ltd" in extract_document_text(".pdf", _pdf_bytes())


def test_document_text_extraction_rejects_scanned_or_empty_pdf(monkeypatch):
    monkeypatch.setattr(
        "app.application.documents.imports.get_settings",
        lambda: SimpleNamespace(
            document_ocr_enabled=False,
            document_ocr_timeout_seconds=30,
            document_ocr_batch_size=8,
        ),
    )
    output = BytesIO()
    pdf = canvas.Canvas(output)
    pdf.showPage()
    pdf.save()
    with pytest.raises(DomainError) as error:
        extract_document_text(".pdf", output.getvalue())
    assert error.value.code == "DOCUMENT_OCR_DISABLED"


def test_document_upload_limit_accepts_exactly_four_mib():
    content = b"a" * (4 * 1024 * 1024)
    assert validate_document_upload("career.txt", content, "text/plain") == (".txt", "text/plain")


def test_document_upload_limit_rejects_one_byte_over_four_mib():
    content = b"a" * (4 * 1024 * 1024 + 1)
    with pytest.raises(DomainError) as error:
        validate_document_upload("career.txt", content, "text/plain")
    assert error.value.code == "DOCUMENT_TOO_LARGE"
    assert error.value.status_code == 413


def _blank_pdf_bytes():
    output = BytesIO()
    pdf = canvas.Canvas(output)
    pdf.showPage()
    pdf.save()
    return output.getvalue()


def test_scanned_pdf_uses_ocr_only_when_explicitly_enabled(monkeypatch):
    calls = []
    monkeypatch.setattr(
        "app.application.documents.imports.get_settings",
        lambda: SimpleNamespace(
            document_ocr_enabled=True,
            document_ocr_timeout_seconds=19,
            document_ocr_batch_size=3,
        ),
    )

    def fake_ocr(content, page_indices, *, timeout_seconds, batch_size):
        calls.append((content, page_indices, timeout_seconds, batch_size))
        return {0: "OCR candidate text"}

    monkeypatch.setattr("app.application.documents.imports.ocr_pdf_pages", fake_ocr)
    text = extract_document_text(".pdf", _blank_pdf_bytes())

    assert text == "OCR candidate text"
    assert len(calls) == 1
    assert calls[0][1:] == ([0], 19, 3)


@pytest.mark.parametrize(
    ("provider_error", "expected_code", "expected_status"),
    [
        ("unavailable", "DOCUMENT_OCR_UNAVAILABLE", 503),
        ("failed", "DOCUMENT_OCR_FAILED", 502),
    ],
)
def test_scanned_pdf_maps_ocr_provider_failures_to_safe_domain_errors(
    monkeypatch, provider_error, expected_code, expected_status
):
    from app.infrastructure.ocr.google_vision import OCRProviderError, OCRUnavailableError

    monkeypatch.setattr(
        "app.application.documents.imports.get_settings",
        lambda: SimpleNamespace(
            document_ocr_enabled=True,
            document_ocr_timeout_seconds=30,
            document_ocr_batch_size=8,
        ),
    )

    def fail_ocr(*_args, **_kwargs):
        if provider_error == "unavailable":
            raise OCRUnavailableError("secret credentials/provider detail")
        raise OCRProviderError("secret upstream response")

    monkeypatch.setattr("app.application.documents.imports.ocr_pdf_pages", fail_ocr)
    with pytest.raises(DomainError) as error:
        extract_document_text(".pdf", _blank_pdf_bytes())

    assert error.value.code == expected_code
    assert error.value.status_code == expected_status
    assert "secret" not in str(error.value)


def test_scanned_pdf_rejects_ocr_response_without_readable_text(monkeypatch):
    monkeypatch.setattr(
        "app.application.documents.imports.get_settings",
        lambda: SimpleNamespace(
            document_ocr_enabled=True,
            document_ocr_timeout_seconds=30,
            document_ocr_batch_size=8,
        ),
    )
    monkeypatch.setattr(
        "app.application.documents.imports.ocr_pdf_pages",
        lambda *_args, **_kwargs: {0: "  "},
    )
    with pytest.raises(DomainError) as error:
        extract_document_text(".pdf", _blank_pdf_bytes())

    assert error.value.code == "DOCUMENT_OCR_TEXT_NOT_FOUND"
    assert error.value.status_code == 422


def test_docx_validation_rejects_zip_without_required_word_parts():
    output = BytesIO()
    import zipfile

    with zipfile.ZipFile(output, "w") as archive:
        archive.writestr("not-a-document.txt", "no Word document parts")

    with pytest.raises(DomainError) as error:
        validate_document_upload(
            "career.docx",
            output.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

    assert error.value.code == "DOCUMENT_SIGNATURE_INVALID"


def test_pdf_page_limit_is_enforced_before_text_extraction():
    output = BytesIO()
    pdf = canvas.Canvas(output)
    from app.application.documents.imports import MAX_PDF_PAGES

    for _ in range(MAX_PDF_PAGES + 1):
        pdf.showPage()
    pdf.save()

    with pytest.raises(DomainError) as error:
        extract_document_text(".pdf", output.getvalue())

    assert error.value.code == "DOCUMENT_PAGE_LIMIT_EXCEEDED"
    assert error.value.status_code == 413
